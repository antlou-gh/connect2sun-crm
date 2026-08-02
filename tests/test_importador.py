"""Testes ao importador de Excel do módulo financeiro.

É a lógica mais complicada da app e a que mexe em dados financeiros: lê uma
folha externa, converte tipos, resolve clientes por NIF e escreve na tabela
`transacoes`. Um erro aqui corrompe contabilidade em silêncio — a importação
diz "importados: 280" na mesma.

Os testes constroem ficheiros .xlsx pequenos em vez de dependerem do
`BD Finanças_2026.xlsx` real, que não está (nem deve estar) no repositório.
"""

from datetime import date

import pytest
from openpyxl import Workbook

from app.financeiro_service import (
    CABECALHOS_IMPORT,
    importar_financeiro,
    mes_para_numero,
    normalizar_nif,
)
from app.models import Transacao

from .conftest import make_cliente

# Uma linha completa e válida, para servir de base às variações.
LINHA_BASE = {
    "Nº de ordem": 1,
    "Descrição": "Painéis solares",
    "NIF": 123456789,
    "Valor total": 1230.0,
    "Entidade emissora": "Fornecedor X",
    "Nº de factura": "FT 2026/1",
    "S/ IVA": 1000.0,
    "IVA %": 23.0,
    "Dia": 15,
    "Mês": "Março",
    "Estado": "Fechado",
    "Tipo de movimento": "Material/Serviços",
}


def _criar_xlsx(tmp_path, linhas, cabecalhos=None, nome="teste.xlsx"):
    """Constrói um .xlsx com cabeçalhos na linha 2 e os dados a partir da 3.

    Sem a Tabela2 definida, o importador cai no fallback documentado
    (cabeçalhos na linha 2, colunas B..O) — que é o que reproduzimos aqui.
    """
    cabecalhos = cabecalhos or list(CABECALHOS_IMPORT)
    wb = Workbook()
    ws = wb.active
    for i, cab in enumerate(cabecalhos, start=2):  # começa na coluna B
        ws.cell(row=2, column=i, value=cab)
    for j, linha in enumerate(linhas, start=3):
        for i, cab in enumerate(cabecalhos, start=2):
            ws.cell(row=j, column=i, value=linha.get(cab))
    caminho = tmp_path / nome
    wb.save(caminho)
    return str(caminho)


# ── normalizar_nif ────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("entrada", "esperado"),
    [
        (123456789, "123456789"),
        (123456789.0, "123456789"),  # o Excel devolve números como float
        ("123456789", "123456789"),
        ("  123 456 789 ", "123456789"),
        ("PT123456789", "123456789"),
        ("pt123456789", "123456789"),
        (None, None),
        ("", None),
        ("   ", None),
    ],
)
def test_normalizar_nif(entrada, esperado):
    assert normalizar_nif(entrada) == esperado


# ── mes_para_numero ───────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("entrada", "esperado"),
    [
        ("Janeiro", 1),
        ("janeiro", 1),
        ("  Dezembro ", 12),
        ("Março", 3),
        ("Marco", 3),  # sem cedilha
        ("MARÇO", 3),
        ("Mar", 3),  # abreviatura
        ("Set", 9),
        ("Mai", 5),
        (None, None),
        ("", None),
        ("Smarch", None),
    ],
)
def test_mes_para_numero(entrada, esperado):
    assert mes_para_numero(entrada) == esperado


# ── Importação: caminho feliz ─────────────────────────────────────────────────


def test_importa_linha_completa(app, db, tmp_path):
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])
    resumo = importar_financeiro(caminho, ano=2026)

    assert resumo["importados"] == 1
    assert resumo["ignoradas"] == 0

    t = Transacao.query.one()
    assert t.numero_ordem == 1
    assert t.descricao == "Painéis solares"
    assert t.valor == 1230.0
    assert t.valor_siva == 1000.0
    assert t.data == date(2026, 3, 15)
    assert t.estado == "Fechado"
    assert t.tipo_movimento == "Material/Serviços"
    assert t.num_factura == "FT 2026/1"


def test_iva_e_calculado_e_nao_lido(app, db, tmp_path):
    """Na folha o IVA é uma fórmula; guardamos o valor, não a string."""
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])
    importar_financeiro(caminho, ano=2026)
    assert Transacao.query.one().iva == 230.0


def test_iva_usa_o_valor_absoluto_dos_dois(app, db, tmp_path):
    """Se total e S/ IVA vierem ambos negativos, a subtracção não pode virar soma.

    Aconteceu em produção (ordem 141, "Pag. acessórios"): valor −26,44 e
    S/ IVA −26,44 gravaram IVA 52,88 em vez de 0.
    """
    linha = {**LINHA_BASE, "Valor total": -1230.0, "S/ IVA": -1000.0}
    caminho = _criar_xlsx(tmp_path, [linha])
    importar_financeiro(caminho, ano=2026)
    assert Transacao.query.one().iva == 230.0


def test_iva_negativo_gera_aviso(app, db, tmp_path):
    """Base maior que o total: acontece com retenção na fonte (royalties).

    O valor pago é líquido de retenção, logo fica abaixo da base e a diferença
    deixa de ser IVA. A linha entra na mesma, mas com aviso — não pode aterrar
    em silêncio na contabilidade.
    """
    linha = {**LINHA_BASE, "Valor total": -1764.85, "S/ IVA": 2215.07}
    caminho = _criar_xlsx(tmp_path, [linha])
    resumo = importar_financeiro(caminho, ano=2026)

    assert resumo["importados"] == 1
    assert len(resumo["avisos"]) == 1
    assert "negativo" in resumo["avisos"][0]
    assert Transacao.query.one().iva < 0


def test_categoria_fica_sempre_por_preencher(app, db, tmp_path):
    """A categorização é feita depois, na app — a importação nunca a inventa."""
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])
    resumo = importar_financeiro(caminho, ano=2026)
    assert Transacao.query.one().categoria is None
    assert resumo["sem_categoria"] == 1


def test_cabecalhos_lidos_por_nome_nao_por_posicao(app, db, tmp_path):
    """Reordenar as colunas na folha não pode partir a importação."""
    baralhados = list(reversed(CABECALHOS_IMPORT))
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE], cabecalhos=baralhados)
    importar_financeiro(caminho, ano=2026)
    t = Transacao.query.one()
    assert t.descricao == "Painéis solares"
    assert t.valor == 1230.0
    assert t.data == date(2026, 3, 15)


# ── Resolução de cliente por NIF ──────────────────────────────────────────────


def test_associa_cliente_pelo_nif(app, db, tmp_path):
    cliente = make_cliente(db, nif="123456789")
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])
    resumo = importar_financeiro(caminho, ano=2026)
    assert Transacao.query.one().cliente_id == cliente.id
    assert resumo["sem_cliente"] == 0


def test_associa_mesmo_com_nif_formatado_de_outra_forma(app, db, tmp_path):
    """O NIF do cliente pode estar gravado com prefixo ou espaços."""
    cliente = make_cliente(db, nif="PT 123 456 789")
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])
    importar_financeiro(caminho, ano=2026)
    assert Transacao.query.one().cliente_id == cliente.id


def test_nif_sem_correspondencia_fica_por_associar(app, db, tmp_path):
    make_cliente(db, nif="999999999")
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])
    resumo = importar_financeiro(caminho, ano=2026)
    assert Transacao.query.one().cliente_id is None
    assert resumo["sem_cliente"] == 1


# ── Idempotência ──────────────────────────────────────────────────────────────


def test_importar_duas_vezes_nao_duplica(app, db, tmp_path):
    """O caso mais perigoso: correr o comando outra vez por engano."""
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE])

    primeiro = importar_financeiro(caminho, ano=2026)
    segundo = importar_financeiro(caminho, ano=2026)

    assert primeiro["importados"] == 1
    assert segundo["importados"] == 0
    assert segundo["ja_existentes"] == 1
    assert Transacao.query.count() == 1


def test_importa_so_as_linhas_novas(app, db, tmp_path):
    """Folha actualizada com linhas novas: importa as novas, salta as antigas."""
    importar_financeiro(_criar_xlsx(tmp_path, [LINHA_BASE], nome="v1.xlsx"), ano=2026)

    nova = {**LINHA_BASE, "Nº de ordem": 2, "Descrição": "Inversor"}
    resumo = importar_financeiro(
        _criar_xlsx(tmp_path, [LINHA_BASE, nova], nome="v2.xlsx"), ano=2026
    )

    assert resumo["importados"] == 1
    assert resumo["ja_existentes"] == 1
    assert Transacao.query.count() == 2


# ── Linhas problemáticas ──────────────────────────────────────────────────────


def test_linha_sem_descricao_e_saltada_em_silencio(app, db, tmp_path):
    """A folha real tem nºs de ordem pré-preenchidos em linhas vazias."""
    fantasma = {"Nº de ordem": 2}
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE, fantasma])
    resumo = importar_financeiro(caminho, ano=2026)

    assert resumo["importados"] == 1
    assert resumo["ignoradas"] == 0  # não conta como problema
    assert resumo["avisos"] == []


@pytest.mark.parametrize(
    ("campo", "valor_mau"),
    [
        ("Nº de ordem", "abc"),
        ("Nº de ordem", None),
        ("Valor total", "n/d"),
        ("Valor total", None),
        ("Dia", 32),
        ("Dia", "quinze"),
        ("Mês", "Smarch"),
        ("Mês", None),
    ],
)
def test_linha_invalida_e_ignorada_com_aviso(app, db, tmp_path, campo, valor_mau):
    """Uma linha má não pode entrar em silêncio nem abortar a importação toda."""
    ma = {**LINHA_BASE, "Nº de ordem": 2, campo: valor_mau}
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE, ma])
    resumo = importar_financeiro(caminho, ano=2026)

    assert resumo["importados"] == 1, "a linha boa devia entrar na mesma"
    assert resumo["ignoradas"] == 1
    assert len(resumo["avisos"]) == 1
    assert Transacao.query.count() == 1


def test_siva_invalido_nao_ignora_a_linha(app, db, tmp_path):
    """S/ IVA é opcional: sem ele a linha entra, com valor_siva e iva a NULL."""
    linha = {**LINHA_BASE, "S/ IVA": "n/d"}
    caminho = _criar_xlsx(tmp_path, [linha])
    resumo = importar_financeiro(caminho, ano=2026)

    assert resumo["importados"] == 1
    t = Transacao.query.one()
    assert t.valor_siva is None
    assert t.iva is None


def test_cabecalho_em_falta_aborta_com_erro_claro(app, db, tmp_path):
    """Melhor rebentar do que importar meia folha com colunas trocadas."""
    incompletos = [c for c in CABECALHOS_IMPORT if c != "NIF"]
    caminho = _criar_xlsx(tmp_path, [LINHA_BASE], cabecalhos=incompletos)

    with pytest.raises(ValueError, match="NIF"):
        importar_financeiro(caminho, ano=2026)

    assert Transacao.query.count() == 0
